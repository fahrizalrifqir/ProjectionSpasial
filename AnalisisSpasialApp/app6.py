import streamlit as st
import pdfplumber
import fitz  # PyMuPDF
import pytesseract
from pdf2image import convert_from_path
import tempfile, os
from PIL import Image
import docx  # python-docx
from docx.shared import RGBColor

# ===============================
# Kriteria resmi per jenis dokumen
# ===============================
CRITERIA = {
    "DELH/DPHL": [
        "Kata Pengantar",
        "Berita Acara Penilaian",
        "Surat Pernyataan Pengelolaan",
        "Surat Pernyataan Ketua Tim",
        "Daftar Isi",
        "Lampiran SK Menteri",
        "RKL-RPL"
    ],
    "KA": [
        "Berita Acara Rapat Tim Teknis Formulir Kerangka Acuan",
        "Kata Pengantar",
        "Berita Acara Rapat Tim Teknis Lanjutan",
        "Saran, Masukan dan Tanggapan Tim Penilai",
        "Surat Pernyataan Pemrakarsa",
        "Peta Tapak Proyek",
        "Peta Batas Ekologis",
        "Peta Batas Sosial",
        "Peta Batas Administrasi",
        "Hasil Konsultasi Publik",
        "Surat Pengantar Penyampaian Dokumen Final KA",
        "Sertifikat Kompetensi Penyusun",
        "Surat Pernyataan Tenaga Ahli"
    ],
    "ANDAL RKL-RPL": [
        "Berita Acara Rapat Tim Teknis",
        "Berita Acara Rapat Komisi",
        "Kata Pengantar",
        "Saran, Masukan dan Tanggapan Tim Penilai",
        "Surat Pernyataan Pemrakarsa",
        "Peta Tapak Proyek",
        "Peta Batas Ekologis",
        "Peta Batas Sosial",
        "Peta Batas Administrasi",
        "Surat Pengantar Penyampaian Dokumen Final ANDAL RKL-RPL",
        "Sertifikat Kompetensi Penyusun",
        "Surat Pernyataan Tenaga Ahli",
        "Surat Pernyataan Kesanggupan"
    ],
    "Addendum ANDAL RKL-RPL": [
        "Berita Acara Rapat Tim Teknis",
        "Berita Acara Rapat Komisi",
        "Kata Pengantar",
        "Saran, Masukan dan Tanggapan Tim Penilai",
        "Peta Tapak Proyek",
        "Surat Pengantar Penyampaian Dokumen Final Addendum",
        "Sertifikat Kompetensi Penyusun",
        "Surat Pernyataan Tenaga Ahli",
        "Surat Pernyataan Kesanggupan"
    ],
    "UKL-UPL": [
        "Berita Acara Rapat Koordinasi",
        "Kata Pengantar",
        "Saran, Masukan dan Tanggapan Tim Penilai",
        "Peta Tapak Proyek",
        "Peta Pengelolaan Lingkungan Hidup",
        "Peta Pemantauan Lingkungan Hidup",
        "Surat Pengantar Penyampaian Dokumen Final UKL-UPL",
        "Surat Pernyataan Kesanggupan",
        "Surat Pernyataan Pemrakarsa"
    ]
}

# ===============================
# Helper untuk cek di Word
# ===============================
def analyze_docx(input_docx, keywords):
    doc = docx.Document(input_docx)
    found = {k: [] for k in keywords}

    for i, para in enumerate(doc.paragraphs):
        for keyword in keywords:
            if keyword.lower() in para.text.lower():
                found[keyword].append(i+1)  # simpan nomor paragraf
                # highlight teks
                run = para.add_run(" [FOUND]")
                run.font.color.rgb = RGBColor(255, 0, 0)

    output_path = input_docx.replace(".docx", "_checked.docx")
    doc.save(output_path)
    return found, output_path

# ===============================
# Streamlit UI
# ===============================
st.set_page_config(page_title="Cek Kelengkapan Dokumen", layout="wide")
st.title("📑 Cek Kelengkapan Dokumen Lingkungan Hidup")

jenis = st.selectbox("Pilih jenis dokumen", list(CRITERIA.keys()))
uploaded_file = st.file_uploader("Upload file PDF atau Word", type=["pdf","docx"])

if uploaded_file:
    KEYWORDS = CRITERIA[jenis]

    # ====================
    # Jika file Word (DOCX)
    # ====================
    if uploaded_file.name.lower().endswith(".docx"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp_input:
            tmp_input.write(uploaded_file.read())
            input_docx = tmp_input.name

        found_pages, output_docx = analyze_docx(input_docx, KEYWORDS)

        st.subheader("📊 Hasil Pengecekan (Word)")
        for keyword, paras in found_pages.items():
            if paras:
                st.success(f"✅ {keyword} ditemukan di paragraf {paras}")
            else:
                st.error(f"❌ {keyword} tidak ditemukan")

        with open(output_docx, "rb") as f:
            st.download_button("⬇️ Download Word hasil cek", f, file_name="dokumen_cek.docx")

        os.remove(input_docx)
        os.remove(output_docx)

    # ====================
    # Jika file PDF
    # ====================
    elif uploaded_file.name.lower().endswith(".pdf"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_input:
            tmp_input.write(uploaded_file.read())
            input_pdf = tmp_input.name

        output_pdf = input_pdf.replace(".pdf", "_checked.pdf")
        found_pages = {k: [] for k in KEYWORDS}

        with pdfplumber.open(input_pdf) as pdf:
            for i, page in enumerate(pdf.pages):
                text = page.extract_text()
                if not text:  # OCR jika halaman berupa gambar
                    images = convert_from_path(input_pdf, first_page=i+1, last_page=i+1)
                    text = pytesseract.image_to_string(images[0], lang="ind")
                for keyword in KEYWORDS:
                    if keyword.lower() in text.lower():
                        found_pages[keyword].append(i+1)

        # Highlight hasil di PDF
        doc = fitz.open(input_pdf)
        for keyword, pages in found_pages.items():
            for page_num in pages:
                page = doc[page_num - 1]
                text_instances = page.search_for(keyword)
                for inst in text_instances:
                    highlight = page.add_highlight_annot(inst)
                    highlight.update()
        doc.save(output_pdf)

        st.subheader("📊 Hasil Pengecekan (PDF)")
        for keyword, pages in found_pages.items():
            if pages:
                st.success(f"✅ {keyword} ditemukan di halaman {pages}")
            else:
                st.error(f"❌ {keyword} tidak ditemukan")

        # Preview halaman PDF hasil
        st.subheader("👀 Preview PDF hasil cek")
        preview_pages = st.slider("Pilih halaman", 1, len(doc), 1)
        page = doc[preview_pages - 1]
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        st.image(img, caption=f"Halaman {preview_pages}", use_container_width=True)

        with open(output_pdf, "rb") as f:
            st.download_button("⬇️ Download PDF hasil cek", f, file_name="dokumen_cek.pdf")

        doc.close()
        os.remove(input_pdf)
        os.remove(output_pdf)
